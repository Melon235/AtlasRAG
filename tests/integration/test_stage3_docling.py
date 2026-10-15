# mypy: disable-error-code="import-not-found, import-untyped, unused-ignore"
"""Real Docling Slim Office integration against deterministic fixtures."""

from __future__ import annotations

from pathlib import Path

import pytest

from atlasrag.domain.errors import ProviderResponseError
from atlasrag.domain.ingestion.enums import ElementType, ListKind, ParseOutcome
from atlasrag.domain.ingestion.source import SourceFile
from atlasrag.providers.parsing.docling import DoclingParser
from atlasrag.services.ingestion.identity import build_candidate_context

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "stage3"
PIPELINE_FINGERPRINT = "f" * 64


def prepared_source(root: Path, fixture_name: str) -> SourceFile:
    path = root / fixture_name
    path.write_bytes((FIXTURES / fixture_name).read_bytes())
    return SourceFile.from_knowledge_directory(root, fixture_name)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_real_docling_slim_parses_docx_semantics(tmp_path: Path) -> None:
    source = prepared_source(tmp_path, "simple.docx")
    context = build_candidate_context(source, PIPELINE_FINGERPRINT)

    parsed = await DoclingParser().parse(source, context)

    assert parsed.parse_info.outcome is ParseOutcome.SUCCESS
    assert parsed.parse_info.parser_name == "docling"
    assert parsed.parse_info.parser_version == "2.135.0"
    assert [element.element_type for element in parsed.elements] == [
        ElementType.HEADING,
        ElementType.PARAGRAPH,
        ElementType.LIST,
        ElementType.LIST,
        ElementType.LIST,
        ElementType.LIST,
        ElementType.CODE,
        ElementType.TABLE,
    ]
    assert parsed.metadata.title == "AtlasRAG Office Fixture"
    assert [element.content for element in parsed.elements[2:4]] == [
        "Alpha item",
        "Beta item",
    ]
    assert all(
        element.metadata.list_kind is ListKind.UNORDERED
        and element.metadata.list_depth == 0
        for element in parsed.elements[2:4]
    )
    assert [element.content for element in parsed.elements[4:6]] == [
        "First numbered item",
        "Second numbered item",
    ]
    assert all(
        element.metadata.list_kind is ListKind.ORDERED
        and element.metadata.list_depth == 0
        for element in parsed.elements[4:6]
    )
    assert parsed.elements[6].content == 'def office_fixture():\n    return "docx"'
    assert parsed.elements[6].metadata.code_language == "python"
    table = parsed.elements[7].table_data
    assert table is not None
    assert table.headers == ("Name", "Value")
    assert table.rows == (("alpha\ncontinued", "1"), ("beta", "2"))


@pytest.mark.integration
@pytest.mark.asyncio
async def test_real_docling_slim_parses_xlsx_grid_and_sheet(tmp_path: Path) -> None:
    source = prepared_source(tmp_path, "simple.xlsx")
    context = build_candidate_context(source, PIPELINE_FINGERPRINT)

    parsed = await DoclingParser().parse(source, context)

    assert parsed.parse_info.outcome is ParseOutcome.SUCCESS
    assert [element.element_type for element in parsed.elements] == [ElementType.TABLE]
    assert parsed.metadata.sheet_names == ("Metrics",)
    table = parsed.elements[0].table_data
    assert table is not None
    assert table.headers == ("Name", "Value")
    assert table.rows == (("alpha", "1"), ("beta", "2"))
    assert table.sheet_name == "Metrics"
    assert table.source_anchor is not None
    assert table.source_anchor.sheet_name == "Metrics"
    assert table.source_anchor.cell_range == "A1:B3"


@pytest.mark.integration
@pytest.mark.asyncio
async def test_real_docling_slim_rejects_uncached_xlsx_formula(
    tmp_path: Path,
) -> None:
    from openpyxl import (
        Workbook,
    )

    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "Value"
    worksheet["A2"] = "=1+2"
    workbook.save(path)
    source = SourceFile.from_knowledge_directory(tmp_path, path.name)
    context = build_candidate_context(source, PIPELINE_FINGERPRINT)

    parsed = await DoclingParser().parse(source, context)

    assert parsed.parse_info.outcome is ParseOutcome.FAILED
    assert parsed.parse_info.failure_code == "UNSUPPORTED_XLSX_FORMULA"
    assert parsed.elements == ()


@pytest.mark.integration
@pytest.mark.asyncio
async def test_real_docling_slim_preserves_merged_xlsx_title_anchor(
    tmp_path: Path,
) -> None:
    from openpyxl import (
        Workbook,
    )

    path = tmp_path / "titled.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Metrics"
    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "Quarterly Metrics"
    worksheet.append(("Name", "Value"))
    worksheet.append(("alpha", 1))
    workbook.save(path)
    source = SourceFile.from_knowledge_directory(tmp_path, path.name)
    context = build_candidate_context(source, PIPELINE_FINGERPRINT)

    parsed = await DoclingParser().parse(source, context)

    assert [element.element_type for element in parsed.elements] == [
        ElementType.PARAGRAPH,
        ElementType.TABLE,
    ]
    title = parsed.elements[0]
    assert title.content == "Quarterly Metrics"
    assert title.source_anchor is not None
    assert title.source_anchor.sheet_name == "Metrics"
    assert title.source_anchor.cell_range == "A1:B1"
    table = parsed.elements[1].table_data
    assert table is not None
    assert table.source_anchor is not None
    assert table.source_anchor.cell_range == "A2:B3"


@pytest.mark.integration
@pytest.mark.asyncio
async def test_real_docling_slim_rejects_non_body_docx_content(
    tmp_path: Path,
) -> None:
    from docx import Document

    path = tmp_path / "header.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Confidential"
    document.add_paragraph("Body")
    document.save(str(path))
    source = SourceFile.from_knowledge_directory(tmp_path, path.name)
    context = build_candidate_context(source, PIPELINE_FINGERPRINT)

    with pytest.raises(ProviderResponseError):
        await DoclingParser().parse(source, context)
