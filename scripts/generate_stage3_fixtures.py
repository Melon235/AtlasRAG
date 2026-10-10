#!/usr/bin/env python3
# mypy: disable-error-code="import-not-found, import-untyped, unused-ignore"
"""Generate deterministic Stage 3 parser input fixtures."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import cast
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIRECTORY = ROOT / "tests" / "fixtures" / "stage3"
_ARCHIVE_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
_CORE_PROPERTIES_PATH = "docProps/core.xml"
_CORE_MODIFIED_TAG = "{http://purl.org/dc/terms/}modified"
_FIXED_MODIFIED_AT = "2020-01-01T00:00:00Z"


def _normalize_office_member(name: str, payload: bytes) -> bytes:
    if name != _CORE_PROPERTIES_PATH:
        return payload
    root = ElementTree.fromstring(payload)
    modified = root.find(_CORE_MODIFIED_TAG)
    if modified is None:
        raise ValueError("Office core properties omit modified timestamp")
    modified.text = _FIXED_MODIFIED_AT
    return cast(
        bytes,
        ElementTree.tostring(root, encoding="utf-8", xml_declaration=True),
    )


def _deterministic_office_archive(payload: bytes) -> bytes:
    """Normalize OOXML member order, timestamps, permissions, and compression."""
    source = BytesIO(payload)
    output = BytesIO()
    with (
        ZipFile(source, "r") as archive,
        ZipFile(
            output,
            "w",
            compression=ZIP_DEFLATED,
            compresslevel=9,
        ) as normalized,
    ):
        for name in sorted(archive.namelist()):
            original = archive.getinfo(name)
            member = ZipInfo(name, date_time=_ARCHIVE_TIMESTAMP)
            member.compress_type = ZIP_DEFLATED
            member.create_system = 3
            member.external_attr = (0o40755 if name.endswith("/") else 0o100644) << 16
            member.extra = b""
            member.comment = b""
            normalized.writestr(
                member,
                _normalize_office_member(name, archive.read(original)),
            )
    return output.getvalue()


def _pdf_bytes() -> bytes:
    content = b"""q
0.8 w
72 346 360 72 re S
72 370 m 432 370 l S
240 346 m 240 418 l S
Q
q
0.94 g
72 462 360 92 re f
Q
BT
/F3 26 Tf
156 742 Td
(AtlasRAG PDF Fixture) Tj
ET
BT
/F1 12 Tf
72 680 Td
(Canonical paragraph from MinerU flash.) Tj
0 -18 Td
(This fixture preserves stable source order.) Tj
ET
BT
/F1 12 Tf
84 616 Td
(a. Alpha item) Tj
0 -16 Td
(b. Beta item) Tj
0 -16 Td
(c. Gamma item) Tj
ET
BT
/F2 10 Tf
84 530 Td
(def parse_pdf\\(\\):) Tj
0 -14 Td
(    items = [\"alpha\", \"beta\"]) Tj
0 -14 Td
(    for item in items:) Tj
0 -14 Td
(        print\\(item\\)) Tj
0 -14 Td
(    return items) Tj
ET
BT
/F1 11 Tf
84 398 Td
(Name) Tj
174 0 Td
(Value) Tj
-174 -24 Td
(alpha) Tj
174 0 Td
(1) Tj
-174 -24 Td
(beta) Tj
174 0 Td
(2) Tj
ET
"""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R /F2 5 0 R /F3 6 0 R >> >> "
            b"/Contents 7 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
        b"<< /Length "
        + str(len(content)).encode("ascii")
        + b" >>\nstream\n"
        + content
        + b"endstream",
    ]

    document = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(document))
        document.extend(f"{number} 0 obj\n".encode("ascii"))
        document.extend(body)
        document.extend(b"\nendobj\n")

    xref_offset = len(document)
    document.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    document.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        document.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    document.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(document)


def _docx_bytes() -> bytes:
    from docx import (
        Document,
    )
    from docx.enum.style import (
        WD_STYLE_TYPE,
    )
    from docx.shared import (
        Pt,
    )

    document = Document()
    properties = document.core_properties
    properties.author = "AtlasRAG"
    properties.last_modified_by = "AtlasRAG"
    properties.title = "AtlasRAG Office Fixture"
    properties.created = datetime(2020, 1, 1)
    properties.modified = datetime(2020, 1, 1)

    document.add_heading("AtlasRAG Office Fixture", level=1)
    document.add_paragraph("Canonical paragraph from Docling.")
    document.add_paragraph("Alpha item", style="List Bullet")
    document.add_paragraph("Beta item", style="List Bullet")
    document.add_paragraph("First numbered item", style="List Number")
    document.add_paragraph("Second numbered item", style="List Number")
    code_style = document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Courier New"
    code_style.font.size = Pt(10)
    document.add_paragraph(
        'def office_fixture():\n    return "docx"',
        style="Code",
    )

    table = document.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        (("Name", "Value"), ("alpha", "1"), ("beta", "2")),
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    table.cell(1, 0).add_paragraph("continued")

    payload = BytesIO()
    document.save(payload)
    return _deterministic_office_archive(payload.getvalue())


def _xlsx_bytes() -> bytes:
    from openpyxl import (
        Workbook,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Metrics"
    for row, values in enumerate(
        (("Name", "Value"), ("alpha", 1), ("beta", 2)),
        start=1,
    ):
        for column, value in enumerate(values, start=1):
            worksheet.cell(row=row, column=column, value=value)
    workbook.properties.creator = "AtlasRAG"
    workbook.properties.lastModifiedBy = "AtlasRAG"
    workbook.properties.created = datetime(2020, 1, 1)
    workbook.properties.modified = datetime(2020, 1, 1)

    payload = BytesIO()
    workbook.save(payload)
    return _deterministic_office_archive(payload.getvalue())


def main() -> None:
    FIXTURE_DIRECTORY.mkdir(parents=True, exist_ok=True)
    (FIXTURE_DIRECTORY / "simple.pdf").write_bytes(_pdf_bytes())
    (FIXTURE_DIRECTORY / "simple.docx").write_bytes(_docx_bytes())
    (FIXTURE_DIRECTORY / "simple.xlsx").write_bytes(_xlsx_bytes())


if __name__ == "__main__":
    main()
